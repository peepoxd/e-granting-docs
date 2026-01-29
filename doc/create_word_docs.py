#!/usr/bin/env python3
"""
PM Exam Assistant - Word Document Generator
Converts Markdown files with Mermaid diagrams to formatted Word documents
"""

import re
import base64
import urllib.parse
import urllib.request
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from io import BytesIO


def set_thai_font(run, font_name='TH Sarabun New', size_pt=14):
    """Set Thai font for a run"""
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    # Set font for complex script (Thai)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run._element.rPr.rFonts.set(qn('w:cs'), font_name)


def extract_mermaid_code(markdown_content):
    """Extract Mermaid diagram code from markdown"""
    pattern = r'```mermaid\n(.*?)```'
    match = re.search(pattern, markdown_content, re.DOTALL)
    if match:
        return match.group(1).strip()
    return None


def mermaid_to_image(mermaid_code):
    """Convert Mermaid code to image using Kroki API"""
    # Use Kroki.io service which supports Mermaid
    url = "https://kroki.io/mermaid/png"

    # Prepare request data with proper headers
    data = mermaid_code.encode('utf-8')
    headers = {
        'Content-Type': 'text/plain',
        'User-Agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36'
    }
    req = urllib.request.Request(url, data=data, headers=headers)

    try:
        with urllib.request.urlopen(req) as response:
            return BytesIO(response.read())
    except Exception as e:
        print(f"Error generating diagram: {e}")
        print(f"  Tip: Try using MCP mermaid server or installing mermaid-cli")
        return None


def add_heading(doc, text, level=1, font_size=16):
    """Add formatted heading"""
    heading = doc.add_heading(level=level)
    run = heading.add_run(text)
    set_thai_font(run, size_pt=font_size)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    return heading


def add_paragraph(doc, text, bold=False, font_size=14, alignment=WD_ALIGN_PARAGRAPH.LEFT):
    """Add formatted paragraph"""
    para = doc.add_paragraph()
    para.alignment = alignment
    run = para.add_run(text)
    set_thai_font(run, size_pt=font_size)
    if bold:
        run.bold = True
    return para


def add_table_from_markdown(doc, markdown_table, font_size=13):
    """Convert markdown table to Word table"""
    lines = [line.strip() for line in markdown_table.strip().split('\n') if line.strip()]

    # Parse header and rows
    rows = []
    for i, line in enumerate(lines):
        if i == 1 and '---' in line:  # Skip separator line
            continue
        cells = [cell.strip() for cell in line.split('|')]
        cells = [c for c in cells if c]  # Remove empty cells
        if cells:
            rows.append(cells)

    if not rows:
        return None

    # Create table
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = 'Light Grid Accent 1'

    # Fill table
    for i, row_data in enumerate(rows):
        for j, cell_text in enumerate(row_data):
            cell = table.rows[i].cells[j]
            cell.text = cell_text.replace('**', '')  # Remove markdown bold

            # Format cell text
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_thai_font(run, size_pt=font_size)
                    if i == 0:  # Header row
                        run.bold = True

    return table


def create_doc_1(md_path, output_path):
    """Create Word document for ข้อ 1 - System Design"""
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    doc = Document()

    # Set page margins (1 inch = 2.54 cm)
    sections = doc.sections
    for section in sections:
        section.page_height = Cm(29.7)  # A4 height
        section.page_width = Cm(21.0)   # A4 width
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)

    # Title
    add_heading(doc, 'ระบบสารสนเทศสนับสนุนโครงการ สสส.', level=1, font_size=16)
    add_paragraph(doc, 'ThaiHealth Project Management Information System (TPMIS)',
                  font_size=12, alignment=WD_ALIGN_PARAGRAPH.CENTER)

    # Extract and add Mermaid diagram
    mermaid_code = extract_mermaid_code(content)
    if mermaid_code:
        add_heading(doc, 'System Architecture Diagram', level=2, font_size=14)
        img_data = mermaid_to_image(mermaid_code)
        if img_data:
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run()
            run.add_picture(img_data, width=Cm(15))

    # Components table
    add_heading(doc, 'คำอธิบาย Components หลัก', level=2, font_size=14)
    components_table = """
| Module | หน้าที่ | ผู้ใช้งานหลัก |
|--------|--------|--------------|
| **Proposal Management** | รับ-ประเมินข้อเสนอโครงการ | ภาคี, เจ้าหน้าที่ |
| **Approval Management** | อนุมัติโครงการและงบประมาณ | ผู้อนุมัติ |
| **Financial Management** | เบิกจ่ายเงินตามงวด, ตรวจหลักฐาน | ฝ่ายการเงิน |
| **Delivery Management** | ติดตามผลงาน, ประเมินความก้าวหน้า | เจ้าหน้าที่ |
| **Closure Management** | สรุปผลและปิดโครงการ | เจ้าหน้าที่, ผู้บริหาร |
"""
    add_table_from_markdown(doc, components_table, font_size=12)

    # Key Features
    add_heading(doc, 'Key Features', level=2, font_size=14)
    features = [
        'E-Proposal: ยื่นข้อเสนอออนไลน์ พร้อมแนบเอกสารดิจิทัล',
        'Workflow Automation: กระบวนการอนุมัติอัตโนมัติตามลำดับขั้น',
        'Financial Tracking: ติดตามการเบิกจ่ายแบบ Real-time เชื่อมต่อ GFMIS',
        'Progress Monitoring: Dashboard ติดตามความก้าวหน้าโครงการ',
        'Notification System: แจ้งเตือนผ่าน Email/LINE อัตโนมัติ'
    ]
    for feature in features:
        para = doc.add_paragraph(style='List Bullet')
        run = para.add_run(feature)
        set_thai_font(run, size_pt=12)

    # Technology Stack
    add_heading(doc, 'Technology Stack', level=2, font_size=14)
    tech_table = """
| Layer | Technology |
|-------|------------|
| **Frontend** | React.js, Responsive Web Design |
| **Backend** | Node.js / .NET Core, REST API |
| **Database** | PostgreSQL, Redis Cache |
| **Security** | OAuth 2.0, SSL/TLS, 2FA |
"""
    add_table_from_markdown(doc, tech_table, font_size=12)

    # Benefits
    add_heading(doc, 'Expected Benefits', level=2, font_size=14)
    benefits_table = """
| ด้าน | ประโยชน์ที่คาดหวัง |
|------|-------------------|
| **ประสิทธิภาพ** | ลดเวลาดำเนินการ 40%, ลดเอกสารกระดาษ 80% |
| **ความโปร่งใส** | ติดตามสถานะได้ Real-time, มี Audit Trail |
| **การตัดสินใจ** | Dashboard สรุปภาพรวมสำหรับผู้บริหาร |
"""
    add_table_from_markdown(doc, benefits_table, font_size=12)

    # Save document
    doc.save(output_path)
    print(f"✓ Created: {output_path}")


def create_doc_2(md_path, output_path):
    """Create Word document for ข้อ 2 - Process Workflow"""
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    doc = Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.page_height = Cm(29.7)
        section.page_width = Cm(21.0)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)

    # Title
    add_heading(doc, 'กระบวนการทำงาน: ระบบจัดการคำขอสินเชื่อออนไลน์', level=1, font_size=16)
    add_paragraph(doc, 'Online Loan Application Management System',
                  font_size=12, alignment=WD_ALIGN_PARAGRAPH.CENTER)

    # Objective
    add_heading(doc, 'วัตถุประสงค์', level=2, font_size=14)
    add_paragraph(doc, 'ระบบจัดการคำขอสินเชื่อออนไลน์ พัฒนาเพื่อให้ลูกค้าสามารถยื่นขอสินเชื่อผ่านช่องทางดิจิทัล ลดระยะเวลาดำเนินการจาก 7 วัน เหลือเพียง 1-2 วัน พร้อมระบบตรวจสอบอัตโนมัติ',
                  font_size=12)

    # Process Flow Diagram
    add_heading(doc, 'Process Flow Diagram', level=2, font_size=14)
    mermaid_code = extract_mermaid_code(content)
    if mermaid_code:
        img_data = mermaid_to_image(mermaid_code)
        if img_data:
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run()
            run.add_picture(img_data, width=Cm(16))

    # As-Is vs To-Be
    add_heading(doc, 'As-Is vs To-Be Comparison', level=2, font_size=14)
    comparison_table = """
| หัวข้อ | As-Is (ก่อนพัฒนา) | To-Be (หลังพัฒนา) |
|--------|-------------------|-------------------|
| **ช่องทางยื่นคำขอ** | มาสาขาเท่านั้น | Online 24/7 |
| **ระยะเวลาดำเนินการ** | 5-7 วันทำการ | 1-2 วันทำการ |
| **การตรวจสอบเอกสาร** | Manual 100% | Auto 70% + Manual 30% |
"""
    add_table_from_markdown(doc, comparison_table, font_size=12)

    # Pain Points
    add_heading(doc, 'Pain Points ที่ระบบแก้ไขได้', level=2, font_size=14)
    pain_points = [
        'ลดการเดินทาง: ลูกค้าไม่ต้องมาสาขา ยื่นได้ทุกที่ทุกเวลา',
        'ลดเอกสารซ้ำซ้อน: ดึงข้อมูลจากระบบกลาง ไม่ต้องกรอกซ้ำ',
        'ลดความผิดพลาด: ระบบตรวจสอบอัตโนมัติก่อนส่ง',
        'เพิ่มความโปร่งใส: ลูกค้าติดตามสถานะได้ Real-time'
    ]
    for point in pain_points:
        para = doc.add_paragraph(style='List Bullet')
        run = para.add_run(point)
        set_thai_font(run, size_pt=12)

    # Benefits
    add_heading(doc, 'Key Benefits', level=2, font_size=14)
    benefits_table = """
| ผู้มีส่วนได้ส่วนเสีย | ประโยชน์ที่ได้รับ |
|---------------------|------------------|
| **ลูกค้า** | สะดวก รวดเร็ว ติดตามสถานะได้ตลอด |
| **เจ้าหน้าที่** | ลดงาน Manual, โฟกัส Case ยากเท่านั้น |
| **องค์กร** | ลดต้นทุนดำเนินการ 35%, เพิ่ม Conversion Rate 25% |
"""
    add_table_from_markdown(doc, benefits_table, font_size=12)

    # Save document
    doc.save(output_path)
    print(f"✓ Created: {output_path}")


def main():
    """Main function to generate Word documents"""
    print("PM Exam Assistant - Word Document Generator")
    print("=" * 50)

    # Define file paths
    doc1_md = Path("ข้อ1_ระบบสารสนเทศ_สสส.md")
    doc2_md = Path("ข้อ2_กระบวนการทำงาน_ระบบสินเชื่อ.md")

    doc1_output = Path("ข้อ1_ระบบสารสนเทศ_สสส.docx")
    doc2_output = Path("ข้อ2_กระบวนการทำงาน_ระบบสินเชื่อ.docx")

    # Generate documents
    try:
        if doc1_md.exists():
            print(f"\nGenerating: {doc1_output}")
            create_doc_1(doc1_md, doc1_output)
        else:
            print(f"⚠ File not found: {doc1_md}")

        if doc2_md.exists():
            print(f"\nGenerating: {doc2_output}")
            create_doc_2(doc2_md, doc2_output)
        else:
            print(f"⚠ File not found: {doc2_md}")

        print("\n" + "=" * 50)
        print("✓ All documents generated successfully!")
        print("\nNote: Please verify that TH Sarabun New font is installed")
        print("      and adjust font size (12-14pt) if content overflows 1 page")

    except Exception as e:
        print(f"\n✗ Error: {e}")
        raise


if __name__ == "__main__":
    main()
